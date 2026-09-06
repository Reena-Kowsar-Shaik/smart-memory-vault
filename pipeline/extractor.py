import io
from typing import Dict, Any

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

def extract_from_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """Extracts text and page count from multi-page PDFs."""
    pages_text = []
    page_count = 0

    # Method 1: pdfplumber (best for layout and clean text)
    if pdfplumber is not None:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                page_count = len(pdf.pages)
                for page in pdf.pages:
                    text = page.extract_text() or ""
                    pages_text.append(text)
                full_text = "\n\n".join(pages_text)
                return {
                    "text": full_text,
                    "page_count": page_count,
                    "is_scanned": len(full_text.strip()) == 0
                }
        except Exception:
            pass

    # Method 2: pypdf fallback
    if pypdf is not None:
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    return {"text": "", "page_count": 0, "error": "Password-protected PDF cannot be read."}
            page_count = len(reader.pages)
            for page in reader.pages:
                pages_text.append(page.extract_text() or "")
            full_text = "\n\n".join(pages_text)
            return {
                "text": full_text,
                "page_count": page_count,
                "is_scanned": len(full_text.strip()) == 0
            }
        except Exception as e:
            return {"text": "", "page_count": 0, "error": f"PDF extraction failed: {str(e)}"}

    return {"text": "", "page_count": 0, "error": "No PDF library installed."}

def extract_from_docx(file_bytes: bytes) -> Dict[str, Any]:
    """Extracts text from Word documents (paragraphs and tables)."""
    if docx is None:
        return {"text": "", "page_count": 1, "error": "python-docx library not installed."}
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs.append(" | ".join(row_text))
        return {"text": "\n\n".join(paragraphs), "page_count": 1, "is_scanned": False}
    except Exception as e:
        return {"text": "", "page_count": 0, "error": f"DOCX extraction failed: {str(e)}"}

def extract_from_txt(file_bytes: bytes) -> Dict[str, Any]:
    """Extracts text from plain text and markdown with encoding detection."""
    for enc in ["utf-8", "utf-16", "latin-1", "cp1252"]:
        try:
            return {"text": file_bytes.decode(enc), "page_count": 1, "is_scanned": False}
        except UnicodeDecodeError:
            continue
    return {"text": "", "page_count": 0, "error": "Unable to decode text with standard encodings."}

def extract_text(file_bytes: bytes, file_extension: str) -> Dict[str, Any]:
    """Main routing function to extract text according to file extension."""
    ext = file_extension.lower()
    if ext == ".pdf":
        return extract_from_pdf(file_bytes)
    elif ext == ".docx":
        return extract_from_docx(file_bytes)
    elif ext in [".txt", ".md"]:
        return extract_from_txt(file_bytes)
    else:
        return {"text": "", "page_count": 0, "error": f"Unsupported format: {ext}"}
